# -*- coding: utf-8 -*-
"""批量将配置表 valid 列的 0 改为 1，并同步修正导出的 JSON。

用法:
    python excel_set_valid.py --path <xlsx路径> --sheet <Sheet名> [--json <导出的txt路径>]

- 自动备份 xlsx 为 *.bak
- 同步把 JSON 中 "valid":0 替换为 "valid":1（保持无 BOM UTF-8 紧凑格式）
"""
import argparse
import shutil

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--path", required=True, help="xlsx 文件路径")
    parser.add_argument("--sheet", required=True, help="Sheet 名")
    parser.add_argument("--json", default=None, help="导出的 JsonText txt 路径（可选，同步修正）")
    parser.add_argument("--include-empty", action="store_true",
                        help="valid 为空（None）的行也视为 0 一并改为 1（默认跳过空行）")
    args = parser.parse_args()

    # 1. 备份
    shutil.copy2(args.path, args.path + ".bak")

    # 2. 修改 Excel
    wb = openpyxl.load_workbook(args.path)
    ws = wb[args.sheet]
    headers = {c.value: c.column for c in ws[1]}
    if "valid" not in headers:
        print("错误: 表中没有 valid 列")
        return
    vc = headers["valid"]
    changed_rows = []
    for row_idx in range(4, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=vc)
        if cell.value is None and not args.include_empty:
            continue
        try:
            if cell.value is None or int(cell.value) == 0:
                cell.value = 1
                changed_rows.append(row_idx)
        except (TypeError, ValueError):
            pass
    wb.save(args.path)
    wb.close()
    print(f"Excel 修改完成: {len(changed_rows)} 行 valid 0->1")
    if changed_rows:
        print(f"行号: {changed_rows[0]}~{changed_rows[-1]} 共 {len(changed_rows)} 行")

    # 3. 同步 JSON
    if args.json:
        with open(args.json, "r", encoding="utf-8") as f:
            content = f.read()
        count = content.count('"valid":0')
        content = content.replace('"valid":0', '"valid":1')
        with open(args.json, "w", encoding="utf-8", newline="") as f:
            f.write(content)
        print(f"JSON 修改完成: {count} 处 valid 0->1")


if __name__ == "__main__":
    main()
